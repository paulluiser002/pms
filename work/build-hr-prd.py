from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(__file__).resolve().parents[1]
source = root / 'outputs' / 'HRDotNet-Performance-Product-Requirements.md'
output = source.with_suffix('.docx')
text = source.read_text(encoding='utf-8')
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
section.footer_distance = Inches(.3)
styles = doc.styles
for key in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet']:
    style = styles[key]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.widow_control = True
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(4)
styles['Normal'].paragraph_format.line_spacing = 1.04
styles['Title'].font.size = Pt(25)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(14)
for name, size in [('Heading 1',16), ('Heading 2',12), ('Heading 3',11)]:
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].paragraph_format.space_before = Pt(14)
    styles[name].paragraph_format.space_after = Pt(7)
    styles[name].paragraph_format.keep_with_next = True

def add_runs(p, value):
    # Markdown source deliberately avoids rich inline styling except requirement identifiers.
    match = re.match(r'^([A-Z]+-[A-Z]+-\d{3})\s+(.*)', value)
    if match:
        p.add_run(match.group(1) + ' ').bold = True
        p.add_run(match.group(2))
    else:
        p.add_run(value)

def add_table(lines):
    values = [[c.strip() for c in ln.strip().strip('|').split('|')] for ln in lines]
    values = [row for row in values if not all(re.fullmatch(r':?-+:?', c) for c in row)]
    count = len(values[0])
    table = doc.add_table(rows=0, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = {2:[2.0,4.77],3:[1.5,2.9,2.37],4:[1.75,1.72,1.65,1.65]}[count]
    widths = {'Version':[.7,1.25,4.82],'Role':[1.15,2.85,2.77],'Order':[.5,1.8,4.47],'Record':[1.2,2.2,3.37],'Employee':[1.1,1.45,4.22],'Decision':[1.4,2.2,3.17]}.get(values[0][0],widths)
    for col, width in zip(table.columns,widths): col.width=Inches(width)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        node=OxmlElement('w:'+edge);node.set(qn('w:w'),'90');node.set(qn('w:type'),'dxa');margins.append(node)
    props.append(margins)
    for ri, rowdata in enumerate(values):
        row=table.add_row()
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for ci,(cell,value) in enumerate(zip(row.cells,rowdata)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
            if ri==0 or (len(values)==3 and ri==1):p.paragraph_format.keep_with_next=True
            run=p.add_run(value);run.font.size=Pt(9)
            if ri==0:
                run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'253E55' if ri==0 else ('F2F5F8' if ri%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

logo=doc.add_paragraph();logo.paragraph_format.space_after=Pt(18)
logo.add_run().add_picture(str(root/'performance/dist/hrdotnet-performance-logo.png'),width=Inches(2.7))
logo.runs[0]._r.xpath('.//wp:docPr')[0].set('descr','Official HRDotNet Performance logo')
lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1;continue
    if line.startswith('|'):
        batch=[]
        while i<len(lines) and lines[i].strip().startswith('|'):batch.append(lines[i]);i+=1
        add_table(batch);continue
    if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
    elif line.startswith('- '):add_runs(doc.add_paragraph(style='List Bullet'),line[2:])
    else:add_runs(doc.add_paragraph(),line)
    i+=1

footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('HRDotNet Performance PRD 2.1  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('777777')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='HRDotNet Performance Product Requirements Document'
doc.core_properties.subject='Performance management MVP scope workflows permissions and acceptance criteria'
doc.core_properties.author=''
doc.core_properties.keywords='HRDotNet Performance, PRD, MVP, performance management'
for element in [*doc.styles.element.iter(qn('w:pBdr')), *doc.element.iter(qn('w:pBdr'))]:
    element.getparent().remove(element)
doc.save(output)
print(f'Markdown words: {len(text.split())}')
print(f'DOCX paragraphs: {len(doc.paragraphs)}, tables: {len(doc.tables)}')
print(output)
